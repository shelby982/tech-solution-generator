"""infra/parser 测试：DOCX TOC/标题扫描兜底 + 公共工具单测。"""
import io
import pytest

from infra.parser import parse_document, ParsedDocument
from infra.parser.toc import (
    extract_special_marks,
    detect_level_from_numbering,
    clean_title,
    _clean_toc_line,
    _looks_like_section_number,
    _text_similarity,
)


# ─────────────────────────────────────────────
# toc.py 公共工具
# ─────────────────────────────────────────────

def test_extract_special_marks_strips_and_collects():
    marks, cleaned = extract_special_marks("★▲技术要求")
    assert "★" in marks and "▲" in marks
    assert cleaned == "技术要求"


def test_extract_special_marks_no_marks():
    marks, cleaned = extract_special_marks("普通章节")
    assert marks == []
    assert cleaned == "普通章节"


def test_detect_level_from_numbering_all_levels():
    assert detect_level_from_numbering("一、概述") == 1
    assert detect_level_from_numbering("第三章 系统设计") == 1
    assert detect_level_from_numbering("1. 项目背景") == 2
    assert detect_level_from_numbering("（1）需求分析") == 3
    assert detect_level_from_numbering("(2) 需求分析") == 3
    assert detect_level_from_numbering("①概述") == 4
    assert detect_level_from_numbering("1.1 子项") == 4


def test_detect_level_returns_none_for_non_heading():
    assert detect_level_from_numbering("普通段落正文") is None
    assert detect_level_from_numbering("") is None


def test_looks_like_section_number():
    """十进制编号要求整数部分 ≤2 位；非十进制编号一律放行。"""
    assert _looks_like_section_number("1.1 ")
    assert _looks_like_section_number("12.3 ")
    assert _looks_like_section_number("1.2.3 ")
    assert not _looks_like_section_number("2320.48 ")  # 表格里的金额
    assert not _looks_like_section_number("902.36 ")
    # 「（1）」「①」不是十进制编号，交给各自的 level 规则
    assert _looks_like_section_number("（1）")
    assert _looks_like_section_number("①")
    assert _looks_like_section_number("")


def test_detect_level_rejects_table_quantities_and_list_items():
    """采购文件表格里的数值和正文枚举项被原来的正则吃成标题。"""
    assert detect_level_from_numbering("2320.48 否") is None
    assert detect_level_from_numbering("902.36 否") is None
    # 列表项：以分号/句号收尾的（1）枚举不是标题
    assert detect_level_from_numbering("（1）异议提出人的名称、地址及有效联系方式；") is None
    assert detect_level_from_numbering("（1）") is None
    # 真标题照旧
    assert detect_level_from_numbering("（1）需求分析") == 3
    assert detect_level_from_numbering("5.2 开标时间") == 4


def test_clean_title_collapses_whitespace():
    assert clean_title("  标题   有空格  ") == "标题 有空格"


def test_clean_toc_line_strips_dot_leaders_and_pages():
    assert _clean_toc_line("第一章 概述...........5") == "第一章 概述"
    assert _clean_toc_line("一、技术要求\t\t12") == "一、技术要求"


def test_text_similarity_prefix_and_containment():
    assert _text_similarity("abcdef", "abcxyz") == pytest.approx(3 / 6)
    assert _text_similarity("abc", "abcdefg") == 0.9  # 包含关系
    assert _text_similarity("", "abc") == 0.0


# ─────────────────────────────────────────────
# DOCX 集成：内存构造一个小 docx，走标题扫描兜底路径
# ─────────────────────────────────────────────

def _build_minimal_docx() -> io.BytesIO:
    """构造无 TOC 的最小 docx：两个 Heading 1 + 一个表格。"""
    from docx import Document

    doc = Document()
    doc.add_heading("第一章 项目概述", level=1)
    doc.add_paragraph("这是项目背景介绍段落。")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "指标"
    table.rows[0].cells[1].text = "数值"
    table.rows[1].cells[0].text = "工期"
    table.rows[1].cells[1].text = "60 天"
    doc.add_heading("第二章 技术方案", level=1)
    doc.add_paragraph("这是技术方案段落。")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def test_parse_docx_heading_scan_returns_sections():
    """无 TOC 文档走标题扫描兜底，识别两个章节并把表格并入正文。"""
    buf = _build_minimal_docx()
    parsed = parse_document(buf, suffix=".docx", filename="test.docx")

    assert isinstance(parsed, ParsedDocument)
    assert len(parsed.sections) == 2

    s1 = parsed.sections[0]
    assert "项目概述" in s1.title
    assert "项目背景" in s1.raw_content
    # 表格应被 markdown 化并出现在 s1 的正文里
    assert "| 指标" in s1.raw_content
    assert "工期" in s1.raw_content

    s2 = parsed.sections[1]
    assert "技术方案" in s2.title
    assert "技术方案段落" in s2.raw_content


def test_parse_document_unsupported_suffix_raises():
    """不支持的扩展名抛 ValueError。"""
    with pytest.raises(ValueError):
        parse_document(io.BytesIO(b"x"), suffix=".txt", filename="x.txt")


# ─────────────────────────────────────────────
# PDF 路径：仅做导入冒烟（不构造 PDF fixture）
# ─────────────────────────────────────────────

def test_pdf_module_imports_cleanly():
    from infra.parser import pdf as pdf_module
    assert hasattr(pdf_module, "parse_pdf")


# ─────────────────────────────────────────────
# PDF 页眉/页脚剔除与逐页字号基线
# ─────────────────────────────────────────────

def _line(text: str, *, page_no: int, in_margin: bool = False, size: float = 10.5) -> dict:
    return {
        "text": text, "avg_size": size, "page_no": page_no,
        "page_body_size": size, "rel_top": 0.5, "rel_bottom": 0.5,
        "in_margin": in_margin,
    }


def test_drop_running_furniture_removes_repeated_margin_lines():
    """14pt 的页眉在正文 10.5pt 的文档里会被字号判据吃成标题 —— 按重复性剔除。"""
    from infra.parser.pdf import _drop_running_furniture

    lines = []
    for page in range(4):
        lines.append(_line("第%d页" % (page + 1), page_no=page, in_margin=True, size=14.0))
        lines.append(_line("招标人：广东电网有限责任公司", page_no=page, in_margin=True, size=14.0))
        lines.append(_line("正文内容 %d" % page, page_no=page))
        lines.append(_line("1.%d 章节" % page, page_no=page))

    kept = _drop_running_furniture(lines)

    assert [ln["text"] for ln in kept] == [
        "正文内容 0", "1.0 章节", "正文内容 1", "1.1 章节",
        "正文内容 2", "1.2 章节", "正文内容 3", "1.3 章节",
    ]


def test_drop_running_furniture_keeps_one_off_margin_line():
    """只在一页边距带出现过的行（脚注）不该被误杀。"""
    from infra.parser.pdf import _drop_running_furniture

    lines = [
        _line("① 本表为格式范本。", page_no=0, in_margin=True),
        _line("正文 A", page_no=0),
        _line("正文 B", page_no=1),
    ]

    assert [ln["text"] for ln in _drop_running_furniture(lines)] == [
        "① 本表为格式范本。", "正文 A", "正文 B",
    ]


def test_drop_running_furniture_page_number_forms():
    """页码长相的行位置无关地剔除；光秃秃的数字只在边距带里才算页码。"""
    from infra.parser.pdf import _drop_running_furniture

    lines = [
        _line("第 12 页", page_no=3),          # 位置无关
        _line("- 7 -", page_no=3),
        _line("Page 3 of 20", page_no=3),
        _line("42", page_no=3, in_margin=True),  # 边距带里的裸数字 = 页码
        _line("2026", page_no=3),                # 正文里的裸数字 = 内容
    ]

    assert [ln["text"] for ln in _drop_running_furniture(lines)] == ["2026"]


def test_page_words_to_lines_uses_page_local_body_size():
    """合同范本页整页 14pt：页基线必须是 14，否则满页都会"比正文大"。"""
    from infra.parser.pdf import _page_words_to_lines

    def word(text, top, size):
        return {"text": text, "top": top, "bottom": top + size, "size": size}

    # 正文页：10.5 为主，一个 14pt 标题
    body_page = [
        word("正文", 300, 10.5), word("正文", 320, 10.5),
        word("正文", 340, 10.5), word("标题", 100, 14.0),
    ]
    lines = _page_words_to_lines(body_page, page_no=0, page_height=800)
    assert {ln["page_body_size"] for ln in lines} == {10.5}

    # 合同范本页：整页 14pt
    template_page = [word("条款一", 100, 14.0), word("条款二", 200, 14.0)]
    lines = _page_words_to_lines(template_page, page_no=1, page_height=800)
    assert {ln["page_body_size"] for ln in lines} == {14.0}


def test_words_to_line_marks_margin_band():
    from infra.parser.pdf import _words_to_line

    def word(text, top):
        return {"text": text, "top": top, "bottom": top + 10, "size": 10.5}

    assert _words_to_line([word("页眉", 20)], 0, 800)["in_margin"]          # 顶部 10% 内
    assert _words_to_line([word("页脚", 740)], 0, 800)["in_margin"]        # 底部 10% 内
    assert not _words_to_line([word("正文", 400)], 0, 800)["in_margin"]
    # 页高缺失（部分 PDF 取不到）时不误判为边距带
    assert not _words_to_line([word("正文", 20)], 0, 0)["in_margin"]


# ─────────────────────────────────────────────
# 表格区域：三列评分表不能被拆成"一个字一行"
# ─────────────────────────────────────────────

def test_rows_to_markdown_shapes_rows():
    from infra.parser.toc import rows_to_markdown

    assert rows_to_markdown([["A", "B"], ["1", "2"]]) == (
        "| A | B |\n| --- | --- |\n| 1 | 2 |"
    )
    assert rows_to_markdown([]) == ""


def test_rows_to_markdown_escapes_and_pads():
    """单元格里的换行/竖线要转义，空单元格与缺列都补齐 —— 否则按 | 切列会错位。"""
    from infra.parser.toc import rows_to_markdown

    lines = rows_to_markdown([["甲|乙", None], ["多\n行"]]).splitlines()
    assert lines[0] == "| 甲\\|乙 |   |"
    assert lines[2] == "| 多<br>行 |   |"


def test_words_outside_tables_filters_by_center():
    """表格内的 word 必须剔除，否则同一批文字会在 raw_content 里出现两遍。"""
    from infra.parser.pdf import _words_outside_tables

    def word(text, x, y):
        return {"text": text, "x0": x, "x1": x + 10, "top": y, "bottom": y + 10}

    words = [word("表内", 100, 100), word("表外", 500, 100)]
    bbox = (90, 90, 200, 200)  # (x0, top, x1, bottom)
    assert [w["text"] for w in _words_outside_tables(words, [bbox])] == ["表外"]
    # 没有表格时原样返回，一个都不能少
    assert len(_words_outside_tables(words, [])) == 2


def test_table_line_is_built_as_body_text():
    from infra.parser.pdf import _table_line

    line = _table_line("| （1）x |", page_no=0, page_height=800, bbox=(10, 100, 500, 700))

    assert line["kind"] == "table"
    assert line["text"] == "| （1）x |"
    # 字号恒为 0、in_margin 恒为 False：表格是正文，不该被标题判据或页眉剔除碰到
    assert line["avg_size"] == 0.0
    assert line["in_margin"] is False
    assert line["top"] == 100


def test_table_rows_never_become_sections():
    """表格是正文不是章节，但仍要留在所属章节的 raw_content 里。"""
    from infra.parser.pdf import _pdf_lines_to_sections

    def heading(text):
        return {"text": text, "avg_size": 10.5, "page_no": 0, "in_margin": False}

    def table_row(text):
        return {"text": text, "kind": "table", "avg_size": 0.0,
                "page_body_size": 0.0, "page_no": 0, "in_margin": False}

    lines = [
        heading("一、项目概况"),
        table_row("| （1）深刻理解项目背景 | 8-10分 |"),
        table_row("| （2）基本理解 | 4-7分 |"),
    ]
    sections, _ = _pdf_lines_to_sections(lines, body_size=10.5, doc_title="T")

    assert [s.title for s in sections] == ["一、项目概况"]
    assert "（1）深刻理解项目背景" in sections[0].raw_content
    assert "（2）基本理解" in sections[0].raw_content


def test_table_only_document_falls_back_when_layout_unknown():
    """形状认不出来的表格宁可退回全文兜底，也不要猜出错位的章节。

    这里的两行只有一列内容可用，推不出「序号列之右、描述列之左」的标题列区间。
    """
    from infra.parser.pdf import _pdf_lines_to_sections

    lines = [
        {"text": "| 说明 |", "kind": "table", "rows": [["说明"], ["内容"]],
         "avg_size": 0.0, "page_body_size": 0.0, "page_no": 0, "in_margin": False},
    ]
    sections, _ = _pdf_lines_to_sections(lines, body_size=12.0, doc_title="评分表")

    assert len(sections) == 1
    assert sections[0].title == "评分表"


# ─────────────────────────────────────────────
# 表格型文档 → 章节（整篇是评分表/报价表，没有散文标题可判）
# ─────────────────────────────────────────────

def _table_line_with_rows(rows: list[list[str]], page_no: int = 0) -> dict:
    from infra.parser.pdf import _table_line

    return _table_line("| md |", page_no=page_no, page_height=800,
                       bbox=(10, 100, 500, 700), rows=rows)


def test_table_layout_finds_serial_column_without_header():
    """无表头：序号列取最左符合的那个（分值列也长着整数的样子，取最左才对）。"""
    from infra.parser.pdf import _table_layout

    rows = [
        ["2.2.4（2）", "技术评分标准", "1", "技术条款差异性", "技术条款差异性",
         "优良（7-10分）：投标响应度较高，完全满足招标技术要求", "10"],
        ["", "", "2", "对项目的理解", "对项目的理解正确、深入、全面",
         "（1）深刻理解项目背景、现状、目标、建设范围", "10"],
        ["", "", "3", "技术方案", "技术方案", "1、方案可行性：", "30"],
    ]
    serial, title_cols = _table_layout(rows)

    assert serial == 2
    assert title_cols == [3, 4]


def test_table_layout_prefers_header_row():
    """有表头：按表头名找标题列，不再依赖序号列。"""
    from infra.parser.pdf import _table_layout

    rows = [
        ["条款号", "", "评分因素", "", "", "评分标准", ""],
        ["2.2.4（3）", "投标报价评分标准", "价格分计算方法名称", "", "",
         "（广东）合理均价基准差径靶心法", ""],
        ["", "", "下浮率", "", "", "2%", ""],
    ]
    serial, title_cols = _table_layout(rows)

    assert serial is None
    assert title_cols == [2]


def test_dedupe_title_keeps_shorter_label():
    """「标签 + 展开」两列留短的；逐字重复只留一条。"""
    from infra.parser.pdf import _dedupe_title

    assert _dedupe_title(["对项目的理解", "对项目的理解正确、深入、全面"]) == "对项目的理解"
    assert _dedupe_title(["技术条款差异性", "技术条款差异性"]) == "技术条款差异性"
    # 互不包含时都留：服务团队 / 项目负责人 才是完整的一行标签
    assert _dedupe_title(["服务团队", "项目负责人"]) == "服务团队 项目负责人"
    # PDF 断行空格要抹掉，否则目录里全是「技术条款 差异性」
    assert _dedupe_title(["技术条款 差异性"]) == "技术条款差异性"


def test_table_only_document_derives_sections_from_rows():
    """整篇是表格时按行还原章节：一行一章，行其余列作该章原文。"""
    from infra.parser.pdf import _pdf_lines_to_sections

    lines = [
        _table_line_with_rows([
            ["2.2.4（2）", "技术评分标准", "1", "技术条款差异性", "技术条款差异性",
             "优良（7-10分）：投标响应度较高", "10"],
            ["", "", "2", "对项目的理解", "对项目的理解正确、深入、全面",
             "（1）深刻理解项目背景", "10"],
        ]),
    ]
    sections, _ = _pdf_lines_to_sections(lines, body_size=12.0, doc_title="评分表")

    assert [s.title for s in sections] == ["1 技术条款差异性", "2 对项目的理解"]
    # 该行的评分标准原文落进本章，下游提炼才有依据
    assert "优良（7-10分）" in sections[0].raw_content
    assert "（1）深刻理解项目背景" in sections[1].raw_content


def test_table_header_row_is_not_a_section():
    """表头行本身不是章节。"""
    from infra.parser.pdf import _pdf_lines_to_sections

    lines = [
        _table_line_with_rows([
            ["条款号", "", "评分因素", "", "", "评分标准", ""],
            ["2.2.4（3）", "", "下浮率", "", "", "2%", ""],
        ]),
    ]
    sections, _ = _pdf_lines_to_sections(lines, body_size=12.0, doc_title="报价表")

    assert [s.title for s in sections] == ["下浮率"]


def test_table_continuation_row_joins_previous_section():
    """跨页断开的续行没有序号也没有标题，并进上一章而不是自成章节。"""
    from infra.parser.pdf import _pdf_lines_to_sections

    lines = [
        _table_line_with_rows([
            ["4", "项目管理", "项目管理方案及质量管理方案",
             "（1）具备科学、合理的项目管理方案", "10"],
        ], page_no=0),
        # 续页：本行是上一行被截断的尾部
        _table_line_with_rows([
            ["", "", "（2）基本符合数字化转型顶层设计，方案总体架构设计比较合理", "", ""],
        ], page_no=1),
    ]
    sections, _ = _pdf_lines_to_sections(lines, body_size=12.0, doc_title="评分表")

    assert [s.title for s in sections] == ["4 项目管理"]
    assert "（2）基本符合数字化转型顶层设计" in sections[0].raw_content


def test_table_sections_lose_no_body_text():
    """还原章节不能丢原文：每行的非标题列都要落进某个章节。"""
    from infra.parser.pdf import _pdf_lines_to_sections

    lines = [
        _table_line_with_rows([
            ["1", "对项目的理解", "（1）深刻理解项目背景", "10"],
            ["2", "技术方案", "1、方案可行性：", "30"],
        ]),
    ]
    sections, _ = _pdf_lines_to_sections(lines, body_size=12.0, doc_title="评分表")

    joined = "\n".join(s.raw_content for s in sections)
    for token in ("深刻理解项目背景", "10", "方案可行性", "30"):
        assert token in joined
