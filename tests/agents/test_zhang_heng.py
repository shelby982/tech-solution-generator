"""张衡 agent 测试 — parse / extract 两个阶段。

依赖 infra.llm 的 mock 路由（环境变量 LLM_MODE=mock）。
"""

import io

import pytest

from agents.zhang_heng import SpecParseResult, ZhangHengAgent
from domain.spec import OutlineMatrixRow
from domain.spec import Section as DomainSection


# ─────────────────────────────────────────────
# 辅助
# ─────────────────────────────────────────────

def _fake_configs():
    """provider 返回最小可用配置，触发 mock 路由。"""
    from services.config_store import LLMConfig
    return ([
        LLMConfig(
            provider="openai",
            api_key="sk-fake",
            base_url="https://example.invalid/v1",
            model="gpt-test",
        ),
    ], 0)


def _build_minimal_docx() -> io.BytesIO:
    """构造最小 docx：两个 Heading 1 + 各一段正文，第二章带 ★ 标记。"""
    from docx import Document

    doc = Document()
    doc.add_heading("第一章 项目概述", level=1)
    doc.add_paragraph("这是项目背景介绍段落。")
    doc.add_heading("第二章 ★技术方案要求", level=1)
    doc.add_paragraph("系统须采用微服务架构。")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ─────────────────────────────────────────────
# parse
# ─────────────────────────────────────────────

@pytest.mark.asyncio
async def test_parse_returns_domain_sections(monkeypatch):
    monkeypatch.setenv("LLM_MODE", "mock")
    agent = ZhangHengAgent(configs_provider=_fake_configs)

    result = await agent.parse(
        [(_build_minimal_docx(), ".docx", "test.docx")],
    )

    assert isinstance(result, SpecParseResult)
    assert len(result.toc) == 2
    assert all(isinstance(s, DomainSection) for s in result.toc)
    assert "项目概述" in result.toc[0].title
    # ★ 标记应在 special_marks 中
    assert "★" in result.toc[1].special_marks
    # 摘要已停用（见 test_parse_does_not_call_doc_summary），恒为空串
    assert result.doc_summary == ""


@pytest.mark.asyncio
async def test_parse_does_not_call_doc_summary(monkeypatch):
    """摘要生成已停用：parse 不再调 dispatch_doc_summary。

    推理型模型在 max_tokens=2000 下必然返回空（预算连思考都不够），跑一次只换来
    一条用户无法处置的降级提示。恢复时这条用例会失败 —— 这是有意的：接回调用必须
    同时给它足够的预算，否则接回来的只是那条提示。
    """
    monkeypatch.setenv("LLM_MODE", "mock")
    import infra.llm

    called = []

    async def _spy(*args, **kwargs):
        called.append(1)
        return "不该被调用的摘要"

    monkeypatch.setattr(infra.llm, "dispatch_doc_summary", _spy)
    agent = ZhangHengAgent(configs_provider=_fake_configs)

    result = await agent.parse([(_build_minimal_docx(), ".docx", "test.docx")])

    assert called == []
    assert result.doc_summary == ""
    # 是「停用」不是「失败」，所以不记 error：别在闸门 1 上吓用户
    assert result.doc_summary_error == ""
    assert len(result.toc) == 2


@pytest.mark.asyncio
async def test_parse_merges_multiple_requirement_files(monkeypatch):
    """多份要求文件的目录按上传顺序拼接，id 连续重编号（各文档原生 id 都是 s1..sN，会撞车）。"""
    monkeypatch.setenv("LLM_MODE", "mock")
    agent = ZhangHengAgent(configs_provider=lambda: ([], 0))

    result = await agent.parse([
        (_build_minimal_docx(), ".docx", "招标文件.docx"),
        (_build_minimal_docx(), ".docx", "评分表.docx"),
    ])

    assert len(result.toc) == 4
    assert [s.id for s in result.toc] == ["s1", "s2", "s3", "s4"]
    # 两份文档的标题拼在一起（fixture 两份同名，所以是两个「第一章 项目概述」）
    assert result.doc_title.count("、") == 1


@pytest.mark.asyncio
async def test_parse_without_sources_raises(monkeypatch):
    """没有任何要求文件时明确抛错，而不是静默产出一份空目录。"""
    monkeypatch.setenv("LLM_MODE", "mock")
    agent = ZhangHengAgent(configs_provider=_fake_configs)

    with pytest.raises(ValueError):
        await agent.parse([])


# ─────────────────────────────────────────────
# extract
# ─────────────────────────────────────────────

@pytest.mark.asyncio
async def test_extract_fills_eight_fields_for_each_section(monkeypatch):
    monkeypatch.setenv("LLM_MODE", "mock")
    agent = ZhangHengAgent(configs_provider=_fake_configs)
    toc = [
        DomainSection(id="s1", level=1, title="技术架构", raw_content="系统须采用微服务"),
        DomainSection(id="s2", level=1, title="安全要求", raw_content="须支持 TLS 1.3", special_marks=["★"]),
    ]

    matrix = await agent.extract(toc)

    assert set(matrix.keys()) == {"s1", "s2"}
    for row in matrix.values():
        assert isinstance(row, OutlineMatrixRow)
        # mock fixture 8 字段全部非空
        assert row.requirement
        assert row.key_points
        assert row.constraint_level in ("mandatory", "recommended", "optional")
        assert row.error == ""


@pytest.mark.asyncio
async def test_extract_returns_empty_rows_when_no_configs(monkeypatch):
    """无 LLM 配置：每行 error 字段写"未配置 LLM"，其它字段保持默认空。"""
    monkeypatch.setenv("LLM_MODE", "mock")
    agent = ZhangHengAgent(configs_provider=lambda: ([], 0))
    toc = [DomainSection(id="s1", level=1, title="X", raw_content="Y")]

    matrix = await agent.extract(toc)

    assert matrix["s1"].error == "未配置 LLM"
    assert matrix["s1"].requirement == ""
    assert matrix["s1"].constraint_level == "recommended"


@pytest.mark.asyncio
async def test_extract_propagates_special_marks_into_dispatch(monkeypatch):
    """带 ★ 的章节：special_marks 被传给 dispatch_outline_json，
    本测试不验证 prompt，仅验证 extract 不丢字段。"""
    monkeypatch.setenv("LLM_MODE", "mock")
    agent = ZhangHengAgent(configs_provider=_fake_configs)
    toc = [
        DomainSection(id="s1", level=1, title="否决项 X", raw_content="必须满足", special_marks=["★"]),
    ]

    matrix = await agent.extract(toc)

    assert matrix["s1"].block_id == "s1"
    assert matrix["s1"].title == "否决项 X"


@pytest.mark.asyncio
async def test_extract_handles_dispatch_error_per_section(monkeypatch):
    """模拟 dispatch_outline_json 对某章节回 error 字段，
    extract 应把 error 字段透传到 OutlineMatrixRow.error，不抛。"""
    monkeypatch.setenv("LLM_MODE", "mock")

    async def fake_dispatch(configs, rr_start, sections):
        for s in sections:
            if "失败" in s["title"]:
                yield {
                    "title": s["title"], "error": "全部 API 失败：mock",
                    "requirement": "", "key_points": "", "veto_items": "",
                    "bonus_items": "", "score_items": "", "evidence_required": "",
                    "constraint_level": "recommended", "indicators": "",
                }
            else:
                yield {
                    "title": s["title"], "error": "",
                    "requirement": "ok", "key_points": "ok", "veto_items": "",
                    "bonus_items": "", "score_items": "", "evidence_required": "",
                    "constraint_level": "recommended", "indicators": "",
                }

    # zhang_heng.py 在函数内部 `from infra.llm import dispatch_outline_json`，
    # 因此 patch 源模块 infra.llm 即可。
    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_json", fake_dispatch)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    toc = [
        DomainSection(id="s1", level=1, title="正常章节", raw_content=""),
        DomainSection(id="s2", level=1, title="失败章节", raw_content=""),
    ]

    matrix = await agent.extract(toc)

    assert matrix["s1"].error == ""
    assert matrix["s1"].requirement == "ok"
    assert "失败" in matrix["s2"].error or "mock" in matrix["s2"].error


# ─────────────────────────────────────────────
# draft_outline
# ─────────────────────────────────────────────

def _no_configs():
    return ([], 0)


@pytest.mark.asyncio
async def test_draft_outline_degrades_without_llm_configs():
    """无 LLM 配置时沿用规范书目录，不抛错，error 写明原因。"""
    agent = ZhangHengAgent(configs_provider=_no_configs)
    source = [
        DomainSection(id="s1", level=1, title="技术方案", raw_content="正文",
                      special_marks=["★"]),
    ]

    result = await agent.draft_outline(source)

    assert result.degraded is True
    assert "未配置模型" in result.error
    assert [(s.id, s.level, s.title) for s in result.sections] == [("s1", 1, "技术方案")]
    # 降级路径要原样带上 raw_content / special_marks，extract 仍能正常跑
    assert result.sections[0].raw_content == "正文"
    assert result.sections[0].special_marks == ["★"]


@pytest.mark.asyncio
async def test_draft_outline_degrades_on_empty_source_toc():
    agent = ZhangHengAgent(configs_provider=_fake_configs)
    result = await agent.draft_outline([])
    assert result.degraded is True
    assert result.sections == []


@pytest.mark.asyncio
async def test_draft_outline_degrades_when_dispatch_raises(monkeypatch):
    async def _boom(configs, rr_start, payload):
        raise RuntimeError("全部 API 失败")

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _boom)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    source = [DomainSection(id="s1", level=1, title="技术方案", raw_content="")]
    result = await agent.draft_outline(source)

    assert result.degraded is True
    assert "全部 API 失败" in result.error
    assert [s.title for s in result.sections] == ["技术方案"]


@pytest.mark.asyncio
async def test_draft_outline_degrades_when_nodes_unparsable(monkeypatch):
    async def _garbage(configs, rr_start, payload):
        return {"nodes": ["不是 dict", {"title": "  "}]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _garbage)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    source = [DomainSection(id="s1", level=1, title="技术方案", raw_content="")]
    result = await agent.draft_outline(source)

    assert result.degraded is True
    assert "无法解析" in result.error


@pytest.mark.asyncio
async def test_draft_outline_degrades_when_too_few_top_level(monkeypatch):
    """一级章节 < MIN_TOP_LEVEL_NODES 视为目录不可信，整体降级。"""
    async def _thin(configs, rr_start, payload):
        return {"nodes": [{"level": 1, "title": "唯一一级"}]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _thin)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    source = [DomainSection(id="s1", level=1, title="技术方案", raw_content="")]
    result = await agent.draft_outline(source)

    assert result.degraded is True
    assert "结构不完整" in result.error


@pytest.mark.asyncio
async def test_draft_outline_locates_by_instruction_before_digest(monkeypatch):
    """提炼要求指向哪一部分，进 prompt 的就是那一部分 —— 不是整份文件。

    实测场景：164 页采购文件解析出 300+ 章，用户要求「按照标包2的技术评分要求
    拆分章节」，但整份文件灌进去时封面/招标公告/合同范本先把 8000 字预算占满。
    """
    captured: dict = {}

    async def _capture(configs, rr_start, payload):
        captured.update(payload)
        return {"nodes": [
            {"level": 1, "title": "项目理解"},
            {"level": 1, "title": "技术响应"},
            {"level": 1, "title": "服务保障"},
        ]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _capture)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    source = (
        [DomainSection(id=f"s{i}", level=1, title=f"第{i}页",
                       raw_content="招标人：广东电网有限责任公司")
         for i in range(1, 200)]
        + [DomainSection(id="s200", level=1, title="标包2 技术评分要求",
                         raw_content="评分因素：技术方案完整性、团队经验、实施计划")]
    )

    await agent.draft_outline(source, instruction="按照标包2的技术评分要求拆分章节")

    digest = captured["spec_digest"]
    assert "标包2 技术评分要求" in digest
    assert "评分因素：技术方案完整性" in digest
    # 噪声章节不该跟着进 prompt
    assert "招标人：广东电网有限责任公司" not in digest
    # 这份语料里没有文档锚点，走的是章级定位 → 没定位到「对应部分」，结构不锁死
    assert captured["follow_structure"] is False


def _anchored_source() -> list[DomainSection]:
    """仿真实采购文件：公告/须知在前，标包1/2/3 的评审标准夹在中间，合同范本在后。"""
    return [
        DomainSection(id="s1", level=1, title="招标公告",
                      raw_content="招标人：某某电网有限责任公司\n招标文件（标准文件范本）"),
        DomainSection(id="s2", level=1, title="投标人须知",
                      raw_content="投标人应当具备下列条件…"),
        DomainSection(id="s3", level=1, title="评审标准", raw_content="\n".join([
            "下列评审标准适用的标的/标包：标包1：关键业务场景研究与验证标包",
            "标包1 的商务要求：具备相关资质",
            "下列评审标准适用的标的/标包：标包2：高可靠技术专题研究与验证标包",
            "技术评分标准：技术方案、项目管理、实施方案、交付成果、服务团队",
            "下列评审标准适用的标的/标包：标包3：主数据管理研究及全过程技术管控标包",
            "标包3 的商务要求：具备相关业绩",
        ])),
        DomainSection(id="s4", level=1, title="合同范本", raw_content="合同条款范本"),
    ]


@pytest.mark.asyncio
async def test_draft_outline_keeps_only_selected_segment_body_in_digest(monkeypatch):
    """定位到标包2 时：只有标包2 的正文进 prompt，其余段落只留标题。

    实测一份 164 页采购文件里标包2 的评分标准只占 7.3%，章级检索命中不了它
    （那 34 个章节的标题里一个「标包2」都没有），只能按文档锚点整段切。
    """
    captured: dict = {}

    async def _capture(configs, rr_start, payload):
        captured.update(payload)
        return {"nodes": [{"level": 1, "title": t} for t in ("项目理解", "技术响应", "服务保障")]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _capture)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    await agent.draft_outline(
        _anchored_source(),
        instruction="按照主招标文件中的标包2的技术评分要求的点进行大纲拆分章节",
    )

    # 定位到了用户点名的那一部分 → 目录要按这部分自身的结构拆
    assert captured["follow_structure"] is True

    digest = captured["spec_digest"]
    # 选中段的正文进来了
    assert "技术评分标准：技术方案、项目管理" in digest
    # 其它标包只留标题，正文不进 —— 这就是省 token 的地方
    assert "标包1：关键业务场景研究与验证标包" in digest
    assert "标包1 的商务要求：具备相关资质" not in digest
    assert "标包3 的商务要求：具备相关业绩" not in digest
    # 文档开头的公告/须知整段丢掉（它们连标题都不该占预算）
    assert "招标人：某某电网有限责任公司" not in digest


@pytest.mark.asyncio
async def test_draft_outline_grounds_from_unselected_content(monkeypatch):
    """「未选中」只影响 digest 里给模型看什么，不影响 grounding 从哪取料。

    这是整个设计能成立的关键：标包2 之外的材料（比如另一份没写标包号的
    技术招标文件）仍然能被 BM25 挂回派生的章节，extract 的 8 字段提炼不缺输入。
    """
    async def _tree(configs, rr_start, payload):
        return {"nodes": [
            {"level": 1, "title": "标包1 的商务要求"},
            {"level": 1, "title": "技术响应"},
            {"level": 1, "title": "服务保障"},
        ]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _tree)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    result = await agent.draft_outline(
        _anchored_source(),
        instruction="按照主招标文件中的标包2的技术评分要求的点进行大纲拆分章节",
    )

    by_title = {s.title: s for s in result.sections}
    assert "具备相关资质" in by_title["标包1 的商务要求"].raw_content


@pytest.mark.asyncio
async def test_draft_outline_skips_scope_model_without_scope_token(monkeypatch):
    """要求里没有范围标识（「按评分项逐条拆章」）时不调模型 —— 无从下手。"""
    seen_payload: list = []

    async def _capture_kw(configs, rr_start, payload):
        seen_payload.append(payload)
        return {"files": [], "keywords": ["标包2"]}

    async def _tree(configs, rr_start, payload):
        return {"nodes": [{"level": 1, "title": t} for t in ("A", "B", "C")]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_scope_select_json", _capture_kw)
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _tree)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    await agent.draft_outline(_anchored_source(), instruction="按评分项逐条拆章")

    assert seen_payload == []


@pytest.mark.asyncio
async def test_draft_outline_scope_model_fallback_runs_once(monkeypatch):
    """点了范围却找不到锚点（文档换了个说法）时才兜底一次，且把标题给模型。"""
    seen_payload: list = []

    async def _capture_kw(configs, rr_start, payload):
        seen_payload.append(payload)
        return {"files": [], "keywords": ["标包2"]}

    async def _tree(configs, rr_start, payload):
        return {"nodes": [{"level": 1, "title": t} for t in ("A", "B", "C")]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_scope_select_json", _capture_kw)
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _tree)

    # 文档里只有「第四包」这类别的写法，锚点族认不出
    source = [
        DomainSection(id="s1", level=1, title=f"第{i}页",
                      raw_content="招标人：某某公司") for i in range(20)
    ]
    agent = ZhangHengAgent(configs_provider=_fake_configs)
    await agent.draft_outline(source, instruction="按照标包2的要求拆分章节")

    assert len(seen_payload) == 1
    assert "第0页" in seen_payload[0]["section_titles"]


@pytest.mark.asyncio
async def test_draft_outline_keeps_full_digest_without_instruction(monkeypatch):
    """没写提炼要求时定位退化为全量，行为与改造前一致。"""
    captured: dict = {}

    async def _capture(configs, rr_start, payload):
        captured.update(payload)
        return {"nodes": [
            {"level": 1, "title": "项目理解"},
            {"level": 1, "title": "技术响应"},
            {"level": 1, "title": "服务保障"},
        ]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _capture)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    source = [
        DomainSection(id="s1", level=1, title="甲章节", raw_content="甲正文"),
        DomainSection(id="s2", level=1, title="乙章节", raw_content="乙正文"),
    ]

    await agent.draft_outline(source, instruction="")

    assert "甲章节" in captured["spec_digest"]
    assert "乙章节" in captured["spec_digest"]
    # 没定位到用户点名的部分 → 结构仍由模型自己组织
    assert captured["follow_structure"] is False


@pytest.mark.asyncio
async def test_draft_outline_builds_hierarchical_ids(monkeypatch):
    async def _tree(configs, rr_start, payload):
        return {"nodes": [
            {"level": 1, "title": "项目理解"},
            {"level": 2, "parent": 0, "title": "需求理解"},
            {"level": 1, "title": "技术响应"},
            {"level": 2, "parent": 2, "title": "架构设计"},
            {"level": 3, "parent": 3, "title": "数据层"},
            {"level": 1, "title": "服务保障"},
        ]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _tree)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    source = [DomainSection(id="s1", level=1, title="技术方案", raw_content="正文")]
    result = await agent.draft_outline(source)

    assert result.degraded is False
    assert [(s.id, s.level) for s in result.sections] == [
        ("s1", 1), ("s1.1", 2), ("s2", 1), ("s2.1", 2), ("s2.1.1", 3), ("s3", 1),
    ]


@pytest.mark.asyncio
async def test_draft_outline_grounds_raw_content_and_inherits_marks(monkeypatch):
    """派生章节靠 BM25 挂回规范书原文，raw_content 与 ★/▲ 都从原文章节继承。"""
    async def _tree(configs, rr_start, payload):
        return {"nodes": [
            {"level": 1, "title": "技术方案"},
            {"level": 1, "title": "实施计划"},
            {"level": 1, "title": "售后服务"},
        ]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _tree)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    source = [
        DomainSection(id="s1", level=1, title="技术方案要求",
                      raw_content="系统须采用微服务架构。", special_marks=["★"]),
        DomainSection(id="s2", level=1, title="实施计划要求",
                      raw_content="工期为 90 日历天。", special_marks=[]),
    ]
    result = await agent.draft_outline(source)

    by_title = {s.title: s for s in result.sections}
    assert "微服务" in by_title["技术方案"].raw_content
    assert by_title["技术方案"].special_marks == ["★"]
    assert "90 日历天" in by_title["实施计划"].raw_content


@pytest.mark.asyncio
async def test_draft_outline_passes_instruction_into_prompt(monkeypatch):
    """用户提炼要求必须真的进 prompt，且带最高优先级标注。"""
    seen = {}

    async def _capture(configs, rr_start, payload):
        seen.update(payload)
        return {"nodes": [
            {"level": 1, "title": "A"}, {"level": 1, "title": "B"},
            {"level": 1, "title": "C"},
        ]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _capture)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    source = [DomainSection(id="s1", level=1, title="技术方案", raw_content="正文")]
    await agent.draft_outline(source, instruction="按评分项逐条拆章", doc_summary="项目摘要")

    assert seen["instruction"] == "按评分项逐条拆章"
    assert seen["doc_summary"] == "项目摘要"
    assert "技术方案" in seen["spec_digest"]
    # 本轮不读素材，扩展点保留但恒为空
    assert seen["material_digest"] == ""
    assert seen["previous_outline"] == ""


@pytest.mark.asyncio
async def test_draft_outline_passes_previous_outline_on_redraft(monkeypatch):
    """整版重出时把上一版目录喂回 prompt，迭代才有连续性。"""
    seen = {}

    async def _capture(configs, rr_start, payload):
        seen.update(payload)
        return {"nodes": [
            {"level": 1, "title": "A"}, {"level": 1, "title": "B"},
            {"level": 1, "title": "C"},
        ]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _capture)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    source = [DomainSection(id="s1", level=1, title="技术方案", raw_content="")]
    previous = [
        DomainSection(id="s1", level=1, title="上一版一级", raw_content=""),
        DomainSection(id="s1.1", level=2, title="上一版二级", raw_content=""),
    ]
    await agent.draft_outline(source, previous_toc=previous)

    assert "上一版一级" in seen["previous_outline"]
    assert "  - 上一版二级" in seen["previous_outline"]


@pytest.mark.asyncio
async def test_draft_outline_reports_warnings_in_error_but_not_degraded(monkeypatch):
    """部分节点被截断/重挂属于「成功但有警告」，不能算降级。"""
    async def _messy(configs, rr_start, payload):
        return {"nodes": [
            {"level": 1, "title": "A"},
            {"level": 1, "title": "A"},
            {"level": 1, "title": "B"},
            {"level": 1, "title": "C"},
        ]}

    import infra.llm
    monkeypatch.setattr(infra.llm, "dispatch_outline_draft_json", _messy)

    agent = ZhangHengAgent(configs_provider=_fake_configs)
    source = [DomainSection(id="s1", level=1, title="技术方案", raw_content="")]
    result = await agent.draft_outline(source)

    assert result.degraded is False
    assert "重名" in result.error
    assert len(result.sections) == 3
