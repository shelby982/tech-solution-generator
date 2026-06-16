"""infra/parser/docx.py — DOCX 解析。

策略：
1. 优先从文档中提取目录结构（Word SDT TOC 或手动目录）
2. 将目录条目映射到正文，提取对应区间的原始内容
3. 无法找到目录时回退到标题扫描
"""

import os
import re
import uuid
import logging
from pathlib import Path
from typing import Optional, Callable, Union, BinaryIO

from ._ocr import ocr_image
from .toc import (
    Section, ParsedDocument, TocEntry, Block,
    HEADING_PATTERNS,
    _TOC_STYLES, _MANUAL_TOC_MARKERS, _SPECIAL_MARK_RE,
    extract_special_marks, detect_level_from_numbering, clean_title,
    _toc_style_to_level, _clean_toc_line,
    _normalize_for_match, _extract_numbering_prefix, _text_similarity,
)

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────
# DOCX body blocks（按 XML 顺序统一表达 paragraph / table）
# ─────────────────────────────────────────────

def _table_to_markdown(table) -> str:
    """把 docx Table 序列化为 Markdown 表格。"""
    rows: list[list[str]] = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = "\n".join(p.text for p in cell.paragraphs).strip()
            text = text.replace("|", "\\|").replace("\n", "<br>")
            cells.append(text or " ")
        rows.append(cells)

    if not rows:
        return ""

    n_cols = max(len(r) for r in rows)
    rows = [r + [" "] * (n_cols - len(r)) for r in rows]

    md_lines = ["| " + " | ".join(rows[0]) + " |",
                "| " + " | ".join(["---"] * n_cols) + " |"]
    for r in rows[1:]:
        md_lines.append("| " + " | ".join(r) + " |")
    return "\n".join(md_lines)


def _ocr_paragraph_images(paragraph, document) -> str:
    """提取段落内所有嵌入图片并 OCR，返回拼接文本。无图片返回空串。"""
    from docx.oxml.ns import qn

    p_elem = paragraph._p
    blips = list(p_elem.iter(qn("a:blip")))
    if not blips:
        return ""

    texts: list[str] = []
    for blip in blips:
        rId = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
        if not rId:
            continue
        try:
            image_part = document.part.related_parts.get(rId)
            if image_part is None:
                continue
            ocr = ocr_image(image_part.blob)
            if ocr:
                texts.append(ocr)
        except Exception as e:
            logger.warning(f"段落图片 OCR 跳过：{e}")
    return "\n".join(texts)


def _iter_body_blocks(document) -> list[Block]:
    """按文档 body 的 XML 顺序构建 block 列表，统一覆盖段落与表格。
    递归进入 SDT 容器（让 Word 内置目录的 TOC 段落以 paragraph 形式出现），
    但不递归进入表格内部。
    """
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    from docx.table import Table

    blocks: list[Block] = []
    body = document.element.body

    _SDT_TAGS = {qn("w:sdt"), qn("w:sdtContent")}
    P_TAG = qn("w:p")
    TBL_TAG = qn("w:tbl")

    def _walk(elem):
        for child in elem.iterchildren():
            tag = child.tag
            if tag == P_TAG:
                para = Paragraph(child, document._body)
                text = para.text
                ocr_text = _ocr_paragraph_images(para, document)
                if ocr_text:
                    text = (text + "\n" + ocr_text).strip() if text.strip() else ocr_text
                style_name = (para.style.name or "") if para.style else ""
                blocks.append(Block(kind="paragraph", text=text, style=style_name, para_obj=para))
            elif tag == TBL_TAG:
                table = Table(child, document._body)
                md = _table_to_markdown(table)
                blocks.append(Block(kind="table", text=md))
            elif tag in _SDT_TAGS:
                _walk(child)

    _walk(body)
    return blocks


# ─────────────────────────────────────────────
# TOC 提取：Word 内置 TOC（SDT）
# ─────────────────────────────────────────────

def _extract_toc_from_sdt(document) -> list[TocEntry]:
    """从 Word 内置 TOC（SDT 结构）中提取目录条目。"""
    try:
        from docx.oxml.ns import qn
    except ImportError:
        return []

    body = document.element.body
    entries: list[TocEntry] = []

    for sdt in body.iter(qn("w:sdt")):
        sdt_pr = sdt.find(qn("w:sdtPr"))
        if sdt_pr is None:
            continue
        doc_part_obj = sdt_pr.find(qn("w:docPartObj"))
        if doc_part_obj is not None:
            gallery = doc_part_obj.find(qn("w:docPartGallery"))
            if gallery is not None and "Table of Contents" in (gallery.get(qn("w:val")) or ""):
                sdt_content = sdt.find(qn("w:sdtContent"))
                if sdt_content is not None:
                    entries = _parse_toc_paragraphs_from_xml(sdt_content, document)
                    if entries:
                        return entries

    # 备用：查找 body 级别的 TOC 样式段落（不在 SDT 内的情况）
    from docx.text.paragraph import Paragraph
    for p_elem in body.iterchildren(qn("w:p")):
        para = Paragraph(p_elem, document._body)
        style_name = (para.style.name or "").lower().strip() if para.style else ""
        if style_name in _TOC_STYLES:
            text = para.text.strip()
            if text:
                level = _toc_style_to_level(style_name)
                text = _clean_toc_line(text)
                marks, cleaned = extract_special_marks(text)
                if cleaned:
                    entries.append(TocEntry(level=level, title=cleaned, raw_title=text, special_marks=marks))

    return entries


def _find_sdt_toc_end(blocks: list[Block]) -> int:
    """扫描 blocks，找到最后一个 TOC 样式段落的索引+1。"""
    last_toc_idx = -1
    scan_limit = min(len(blocks), 300)
    for i in range(scan_limit):
        if blocks[i].kind != "paragraph":
            continue
        style_name = (blocks[i].style or "").lower().strip()
        if style_name in _TOC_STYLES:
            last_toc_idx = i
    return last_toc_idx + 1 if last_toc_idx >= 0 else 0


def _parse_toc_paragraphs_from_xml(sdt_content, document) -> list[TocEntry]:
    """从 SDT Content 的 XML 中解析 TOC 段落。"""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    entries: list[TocEntry] = []
    for p_elem in sdt_content.iter(qn("w:p")):
        para = Paragraph(p_elem, document._body)
        style_name = (para.style.name or "").lower().strip() if para.style else ""
        if style_name not in _TOC_STYLES:
            continue
        text = para.text.strip()
        if not text:
            continue
        level = _toc_style_to_level(style_name)
        text = _clean_toc_line(text)
        marks, cleaned = extract_special_marks(text)
        if cleaned:
            entries.append(TocEntry(level=level, title=cleaned, raw_title=text, special_marks=marks))
    return entries


# ─────────────────────────────────────────────
# TOC 提取：手动编写的目录
# ─────────────────────────────────────────────

def _extract_toc_manual(blocks: list[Block]) -> tuple[list[TocEntry], int]:
    """扫描文档前部，查找手动编写的目录区域并提取条目。"""
    toc_start_idx = -1
    scan_limit = min(len(blocks), 80)

    for i in range(scan_limit):
        if blocks[i].kind != "paragraph":
            continue
        text = blocks[i].text.strip().lower()
        if text in _MANUAL_TOC_MARKERS:
            toc_start_idx = i + 1
            break

    if toc_start_idx < 0:
        return [], 0

    entries: list[TocEntry] = []
    blank_count = 0
    toc_end_idx = toc_start_idx

    for i in range(toc_start_idx, min(len(blocks), toc_start_idx + 200)):
        if blocks[i].kind != "paragraph":
            if entries:
                break
            continue
        raw_text = blocks[i].text.strip()

        if not raw_text:
            blank_count += 1
            if blank_count >= 2 and entries:
                break
            continue
        blank_count = 0

        if len(raw_text) > 120 and not detect_level_from_numbering(raw_text):
            break

        cleaned = _clean_toc_line(raw_text)
        if not cleaned:
            continue

        marks, title = extract_special_marks(cleaned)
        level = detect_level_from_numbering(title)

        if level is not None and len(title) < 100:
            entries.append(TocEntry(level=level, title=title, raw_title=raw_text, special_marks=marks))
            toc_end_idx = i + 1
        elif entries and len(title) < 80:
            style_name = (blocks[i].style or "").lower()
            if "heading" in style_name or "标题" in style_name:
                entries.append(TocEntry(level=1, title=title, raw_title=raw_text, special_marks=marks))
                toc_end_idx = i + 1

    return entries, toc_end_idx


# ─────────────────────────────────────────────
# 内容映射：TOC → 正文
# ─────────────────────────────────────────────

def _map_toc_to_body(
    toc_entries: list[TocEntry],
    blocks: list[Block],
    progress_callback: Optional[Callable] = None,
    body_start: int = 0,
) -> list[tuple[TocEntry, str]]:
    """将目录条目映射到正文标题位置，提取该标题到下一个标题之间的内容。"""
    if progress_callback:
        progress_callback("映射目录到正文...", 0, len(toc_entries))

    block_texts = [(i, blocks[i].text.strip()) for i in range(len(blocks))]

    matched_positions: list[int] = []
    last_pos = body_start

    for entry in toc_entries:
        pos = _find_heading_in_body(entry, block_texts, last_pos, blocks)
        matched_positions.append(pos)
        if pos >= 0:
            last_pos = pos + 1

    matched_count = sum(1 for p in matched_positions if p >= 0)
    logger.info(
        f"TOC正文匹配结果：{matched_count}/{len(toc_entries)} 条命中，"
        f"body_start={body_start}，"
        f"命中位置范围：{[p for p in matched_positions if p >= 0][:5]}..."
    )

    results: list[tuple[TocEntry, str]] = []
    for idx, (entry, pos) in enumerate(zip(toc_entries, matched_positions)):
        if progress_callback:
            progress_callback(entry.title, idx + 1, len(toc_entries))

        if pos < 0:
            results.append((entry, ""))
            continue

        next_pos = len(blocks)
        for j in range(idx + 1, len(matched_positions)):
            if matched_positions[j] >= 0:
                next_pos = matched_positions[j]
                break

        content_parts = []
        for j in range(pos + 1, next_pos):
            text = blocks[j].text.strip()
            if text:
                content_parts.append(text)

        raw_content = "\n".join(content_parts)
        results.append((entry, raw_content))

    return results


def _find_heading_in_body(
    entry: TocEntry,
    block_texts: list[tuple[int, str]],
    start_from: int,
    blocks: list[Block],
) -> int:
    """在 block 列表中找到与 TOC 条目匹配的标题位置（仅在 paragraph block 上匹配）。"""
    target = _normalize_for_match(entry.title)
    if not target:
        return -1

    target_prefix = _extract_numbering_prefix(entry.title)

    best_idx = -1
    best_score = 0

    for i, text in block_texts:
        if i < start_from:
            continue
        if blocks[i].kind != "paragraph":
            continue
        if not text or len(text) > 200:
            continue

        normalized = _normalize_for_match(text)

        if target_prefix:
            text_prefix = _extract_numbering_prefix(text)
            if text_prefix and text_prefix == target_prefix:
                score = _text_similarity(target, normalized)
                if score > best_score:
                    best_score = score
                    best_idx = i
                    if score > 0.8:
                        return i

        score = _text_similarity(target, normalized)
        if score > 0.85 and score > best_score:
            best_score = score
            best_idx = i

    return best_idx


# ─────────────────────────────────────────────
# DOCX 解析入口
# ─────────────────────────────────────────────

def parse_docx(
    file_source: Union[str, BinaryIO],
    filename: str = "",
    progress_callback: Optional[Callable[[str, int, int], None]] = None,
) -> ParsedDocument:
    """解析 DOCX：TOC 优先，回退到标题扫描。"""
    from docx import Document

    doc_id = str(uuid.uuid4())

    if progress_callback:
        progress_callback("读取文档...", 0, -1)

    document = Document(file_source)
    blocks = _iter_body_blocks(document)
    n_para = sum(1 for b in blocks if b.kind == "paragraph")
    n_tbl = sum(1 for b in blocks if b.kind == "table")
    logger.info(f"body blocks: {len(blocks)} (paragraphs={n_para}, tables={n_tbl})")

    doc_title = Path(filename).stem if filename else "文档"
    title_scanned = 0
    for b in blocks:
        if b.kind != "paragraph":
            continue
        title_scanned += 1
        if title_scanned > 20:
            break
        style_name = (b.style or "").lower()
        if "title" in style_name and b.text.strip():
            doc_title = clean_title(b.text)
            break

    if progress_callback:
        progress_callback("提取目录结构...", 0, -1)

    toc_entries = _extract_toc_from_sdt(document)
    toc_source = "Word内置目录"
    body_start = 0

    if toc_entries:
        body_start = _find_sdt_toc_end(blocks)

    if not toc_entries:
        toc_entries, body_start = _extract_toc_manual(blocks)
        toc_source = "手动目录"

    logger.info(f"TOC提取路径：{toc_source}，条目数：{len(toc_entries)}，body_start：{body_start}")

    if not toc_entries:
        if progress_callback:
            progress_callback("未发现目录，使用标题扫描...", 0, -1)
        return _parse_docx_by_heading_scan(blocks, doc_id, doc_title, progress_callback)

    logger.info(f"从{toc_source}提取到 {len(toc_entries)} 个章节")
    if progress_callback:
        progress_callback(f"从{toc_source}提取到 {len(toc_entries)} 个章节", 0, len(toc_entries))

    mapped = _map_toc_to_body(toc_entries, blocks, progress_callback, body_start=body_start)

    if progress_callback:
        progress_callback("提取章节原文...", 0, len(mapped))

    sections: list[Section] = []
    for idx, (entry, raw_content) in enumerate(mapped):
        sections.append(Section(
            section_id=f"s{idx + 1}",
            level=entry.level,
            title=clean_title(entry.title),
            raw_content=raw_content,
            special_marks=entry.special_marks,
        ))
        if progress_callback:
            progress_callback(entry.title, idx + 1, len(mapped))

    return ParsedDocument(doc_id, doc_title, sections)


def _parse_docx_by_heading_scan(
    blocks: list[Block],
    doc_id: str,
    doc_title: str,
    progress_callback: Optional[Callable] = None,
) -> ParsedDocument:
    """无 TOC 时按 block 扫描标题。"""
    sections: list[Section] = []

    if progress_callback:
        progress_callback("扫描章节标题...", 0, -1)

    heading_indices: list[tuple[int, int, str, list[str]]] = []
    for i, block in enumerate(blocks):
        if block.kind != "paragraph":
            continue
        raw_text = block.text.strip()
        if not raw_text:
            continue
        marks, text = extract_special_marks(raw_text)
        level = _get_heading_level(block.style or "", text)
        if level is not None:
            heading_indices.append((i, level, text, marks))

    total_headings = len(heading_indices)
    if progress_callback:
        progress_callback(f"发现 {total_headings} 个章节", 0, total_headings)

    for h_idx, (block_idx, level, text, marks) in enumerate(heading_indices):
        if progress_callback:
            progress_callback(text, h_idx + 1, total_headings)

        next_block_idx = (
            heading_indices[h_idx + 1][0]
            if h_idx + 1 < len(heading_indices)
            else len(blocks)
        )

        hint_parts = []
        for j in range(block_idx + 1, next_block_idx):
            next_text = blocks[j].text.strip()
            if next_text:
                hint_parts.append(next_text)

        raw_content = "\n".join(hint_parts)

        sections.append(Section(
            section_id=f"s{h_idx + 1}",
            level=level,
            title=clean_title(text),
            raw_content=raw_content,
            special_marks=marks,
        ))

    if not sections:
        all_text = "\n".join(b.text.strip() for b in blocks if b.text.strip())
        if all_text:
            logger.info(f"未识别到任何标题，启用纯表格/无标题兜底（{len(all_text)} 字符）")
            sections.append(Section(
                section_id="s1",
                level=1,
                title=doc_title or "全文",
                raw_content=all_text,
                special_marks=[],
            ))

    return ParsedDocument(doc_id, doc_title, sections)


def _get_heading_level(style_name: str, text: str) -> Optional[int]:
    s = style_name.lower().strip()
    heading_map = {
        "heading 1": 1, "heading1": 1, "标题 1": 1, "标题1": 1,
        "heading 2": 2, "heading2": 2, "标题 2": 2, "标题2": 2,
        "heading 3": 3, "heading3": 3, "标题 3": 3, "标题3": 3,
        "heading 4": 4, "heading4": 4, "标题 4": 4, "标题4": 4,
        "heading 5": 5, "heading5": 5, "标题 5": 5, "标题5": 5,
        "heading 6": 6, "heading6": 6, "标题 6": 6, "标题6": 6,
    }
    for key, level in heading_map.items():
        if s == key.lower():
            return level

    if len(text) < 40:
        numbering_level = detect_level_from_numbering(text)
        if numbering_level is not None:
            return numbering_level

    return None


# ─────────────────────────────────────────────
# 旧版 .doc → .docx 转换（依赖 LibreOffice）
# ─────────────────────────────────────────────

def parse_doc_via_convert(
    file_source: Union[str, BinaryIO],
    filename: str = "",
    progress_callback: Optional[Callable[[str, int, int], None]] = None,
) -> ParsedDocument:
    import subprocess
    import tempfile
    import shutil

    if progress_callback:
        progress_callback("检测到旧版 .doc 格式，正在转换...", 0, -1)

    out_dir = tempfile.mkdtemp(prefix="tsg_doc_convert_")
    stem = Path(filename).stem if filename else "document"

    try:
        if isinstance(file_source, str):
            doc_path = file_source
        else:
            doc_path = os.path.join(out_dir, f"{stem}.doc")
            with open(doc_path, "wb") as f:
                f.write(file_source.read())

        try:
            result = subprocess.run(
                ["soffice", "--headless", "--convert-to", "docx", doc_path, "--outdir", out_dir],
                capture_output=True,
                text=True,
                timeout=120,
            )
            if result.returncode != 0:
                raise ValueError(f".doc 转换失败：{result.stderr.strip()}")
        except FileNotFoundError:
            raise ValueError(
                "该文件实际为旧版 .doc 格式（非 .docx），"
                "服务器未安装 LibreOffice 无法自动转换。"
                "请先用 Word 或 WPS 将文件「另存为」.docx 格式后重新上传。"
            )
        except subprocess.TimeoutExpired:
            raise ValueError(".doc 转换超时（120s），请尝试上传较小的文件")

        converted = Path(out_dir) / f"{stem}.docx"
        if not converted.exists():
            raise ValueError(
                "该文件实际为旧版 .doc 格式，自动转换失败。"
                "请先用 Word 或 WPS 将文件「另存为」.docx 格式后重新上传。"
            )

        if progress_callback:
            progress_callback("格式转换完成，开始解析...", 0, -1)

        return parse_docx(str(converted), filename, progress_callback)

    finally:
        shutil.rmtree(out_dir, ignore_errors=True)
